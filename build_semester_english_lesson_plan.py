from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("outputs/AiDi+英文教學單元活動設計_一學期完整版.docx")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

FONT = "DFKai-SB"
FONT_EN = "DFKai-SB"
INK = "000000"
TEAL = "FFF2CC"
TEAL_DARK = "000000"
PALE_TEAL = "F2F2F2"
PALE_GOLD = "FFF2CC"
PALE_GRAY = "F4F7F8"
MID_GRAY = "555555"
WHITE = "000000"
GREEN = "DFF3E6"
RED = "FCE8E6"
TABLE_WIDTH = 9360
TABLE_INDENT = 120


sessions = [
    {
        "no": 1,
        "date": "115 年 4 月 21 日",
        "unit": "Hello, New Friends! 認識彼此與英語起點診斷",
        "source": "第一課／學期初診斷",
        "language": "問候、自我介紹、年齡與課堂常用語",
        "vocab": "hello, please, come in, come here, morning, afternoon, evening, night",
        "patterns": "My name is ___. / I’m ___ years old. / Nice to meet you.",
        "context": "認識彼此、共同訂定學習約定、連結屏東與臺東生活經驗",
        "knowledge": "能辨認常見問候語、字母與 1-20 數字，並理解基本自我介紹句型。",
        "practice": "能以口說、指認或字卡排序完成姓名與年齡介紹，並回應同伴。",
        "literacy": "能安全操作互動題卡，了解數位活動是練習與回饋工具。",
        "lead1": "以姓名卡與生活照片建立安全感；示範 Hello、My name is...、I’m... years old，並說明輪流與尊重的規則。",
        "self": "完成字母、數字、生活用語及常見單字的低壓診斷；可用口說、點選或拖曳作答。",
        "group": "四位學生用自我介紹題卡輪流完成姓名、喜好與期待；同伴從卡片中找出一項共同點。",
        "share": "每人介紹一位同伴；若不方便完整口說，可由同伴代讀、本人點頭確認。",
        "lead2": "統整本學期英語學習方式，確認共同約定，記錄各生的起點表現與偏好的回應方式。",
        "digital": "自我介紹題卡、隨機挑人、起點診斷卡",
        "assessment": "診斷性評量：字母、數字、生活用語辨識；觀察是否能完成至少 2 句自我介紹。",
    },
    {
        "no": 2,
        "date": "115 年 4 月 28 日",
        "unit": "Sound and Color 短母音與顏色探索",
        "source": "第一課",
        "language": "短母音 a/e/i/o/u、顏色單字與詢問顏色",
        "vocab": "red, orange, yellow, green, blue, indigo, purple, brown, white, gray, black, pink",
        "patterns": "What color is it? / It is ___.",
        "context": "從天空、草地、彩虹與三原色混合觀察生活中的顏色",
        "knowledge": "能辨識五個短母音的基本音值，並認讀至少 8 個顏色單字。",
        "practice": "能聽辨短母音、將 CVC 單字拆音合音，並以句型回答顏色。",
        "literacy": "能操作顏色辨識與混色活動，依畫面變化修正英文回答。",
        "lead1": "以 dog、pen、pig、hot、sun 示範母音像聲音的橋；再以生活圖片引出顏色與 What color is it?。",
        "self": "在母音拼音器更換中間母音，比較發音；完成顏色聽辨、英文配對與跟讀。",
        "group": "兩位學生各選一個三原色，第三位猜混色結果，第四位用 It is... 完整作答；輪替角色。",
        "share": "各組說出一組混色式與英文答案，其他人用顏色卡判斷是否一致。",
        "lead2": "回顧短母音與 12 色，針對易混淆的 green/gray、blue/purple 進行對比練習。",
        "digital": "母音自由拼音、顏色認識、顏色混合",
        "assessment": "形成性評量：5 個短母音聽辨、顏色辨識與 What color is it? 口語回應。",
    },
    {
        "no": 3,
        "date": "115 年 5 月 5 日",
        "unit": "Numbers and Kind Words 數字與生活用語",
        "source": "第二課",
        "language": "短母音複習、1-20 數字、年齡與禮貌用語",
        "vocab": "one-twenty, please, thank you, I’m sorry, big, small",
        "patterns": "What number is it? / How old are you? / I’m ___ years old.",
        "context": "生活情境中的請求、感謝與道歉；從海岸景觀比較 big/small",
        "knowledge": "能認讀 1-20 英文數字並理解 please、thank you、I’m sorry 的情境功能。",
        "practice": "能依數字提示說出英文，並在生活情境選擇合適的禮貌語。",
        "literacy": "能在數字與情境卡活動中依回饋修正，不以猜測取代讀題。",
        "lead1": "用 CVC 字卡快速複習短母音；以年齡與數量導入 1-20，再示範三種禮貌語的真實情境。",
        "self": "完成數字聽讀配對與年齡句型；閱讀 3 張情境卡，選擇適合的生活用語。",
        "group": "小組抽取『需要幫忙／受到幫助／不小心犯錯』情境，合作排出對話順序並角色演練。",
        "share": "各組演出一個情境，觀眾指出聽到的關鍵語與使用原因。",
        "lead2": "統整數字、年齡與禮貌用語；用快速口頭題檢核 1-20 和情境判斷。",
        "digital": "數字配對、情境卡、隨機挑人",
        "assessment": "形成性評量：數字正確率、年齡句型完整度、禮貌語情境判斷。",
    },
    {
        "no": 4,
        "date": "115 年 5 月 12 日",
        "unit": "I Am, You Are 主詞、Be 動詞與感謝表達",
        "source": "第二課",
        "language": "主詞 I/you/he/she/it/we/they 與 am/is/are；自我介紹",
        "vocab": "I, you, he, she, it, we, they, happy, friend, family, card, flower, gift",
        "patterns": "I am ___. / You are my friend. / My name is ___. / I like/love ___.",
        "context": "母親節、感謝卡與生活中值得感謝的人事物",
        "knowledge": "能理解主詞是句子的主角，並配對 I-am、you/we/they-are、he/she/it-is。",
        "practice": "能完成簡短自我介紹與感謝句，並從句子中找出主詞。",
        "literacy": "能使用數位卡片編排文字，確認內容後再儲存作品。",
        "lead1": "以人物與物品圖卡建立主詞概念，在位值式句型板上示範主詞＋Be 動詞＋描述。",
        "self": "完成主詞與 Be 動詞配對；拖曳字卡排出 I am happy.、You are my friend.。",
        "group": "每組選一位想感謝的人，合作組成 2-3 句卡片內容，再檢查主詞與 Be 動詞。",
        "share": "讀出或展示感謝卡；同伴以『我聽到的主詞是...』回饋。",
        "lead2": "整理 7 個主詞的 Be 動詞搭配，指出常見錯誤 I is、he are，完成離堂配對題。",
        "digital": "主詞配對、句子排序、感謝卡片",
        "assessment": "形成性評量：主詞辨識、Be 動詞搭配、2 句以上有意義的自我或感謝表達。",
    },
    {
        "no": 5,
        "date": "115 年 5 月 19 日",
        "unit": "This or That 近遠物品與文具描述",
        "source": "第二課",
        "language": "this/that、a/an、文具與顏色＋名詞",
        "vocab": "pencil, eraser, pen, marker, book, ruler, bag; red, blue, yellow...",
        "patterns": "What is this/that? / It is a/an ___. / This/That is a ___.",
        "context": "教室物品、遠近觀察與臺東熱氣球的色彩描述",
        "knowledge": "能分辨 this 與 that 的距離概念，並理解 a/an 依發音選用。",
        "practice": "能以顏色＋物品組成名詞片語，回答至少 4 題近遠物品問題。",
        "literacy": "能在圖像問答中先觀察距離與物品，再提交句子。",
        "lead1": "將真實文具放在近處與遠處，示範 What is this/that?；比較 a pencil 與 an eraser。",
        "self": "操作近遠圖卡，選擇 this/that、a/an 與物品字卡，完成完整句。",
        "group": "進行『教室尋物任務』：提示者只能說距離與顏色，回答者用完整句找出物品。",
        "share": "各組展示一題最容易混淆的任務，其他人解釋使用 this 或 that 的理由。",
        "lead2": "統整距離、冠詞與形容詞順序；以 4 題快問快答檢核句型遷移。",
        "digital": "這那問答、物品圖卡、圖卡配對",
        "assessment": "形成性評量：this/that 判斷、a/an 選用、顏色＋物品的詞序與口說清晰度。",
    },
    {
        "no": 6,
        "date": "115 年 5 月 26 日",
        "unit": "Animals and Long A 動物單字與 A 長母音",
        "source": "第三課",
        "language": "長短母音比較、A 長母音 ai/ay/a_e、動物與 Yes/No 問答",
        "vocab": "fox, bear, monkey, elephant, tiger, lion, zebra, mouse, horse, panda, rabbit, goat",
        "patterns": "Is it a/an ___? / Yes, it is. / No, it isn’t. It is a/an ___.",
        "context": "屏東海洋與陸域動物、生物多樣性與安全觀察",
        "knowledge": "能辨識 A 長母音 /eɪ/ 的常見拼法，並認讀 12 個動物單字。",
        "practice": "能比較 cake/cack 類長短音，並使用 Yes/No 句型猜動物。",
        "literacy": "能在猜圖活動中先聽問題、再選答案，並閱讀系統回饋。",
        "lead1": "以 cake/cap、rain/day 比較長短音，標示 ai、ay、a_e；以動物輪廓導入單字與 a/an。",
        "self": "自由替換 A 長母音前後字母並朗讀；完成動物圖卡與英文配對。",
        "group": "『動物線索局』：一人持隱藏動物卡，其餘輪流問 Is it a/an...?；持卡者只用 Yes/No 回應。",
        "share": "各組分享一題猜錯後如何用 No, it isn’t. It is... 修正，強化完整回答。",
        "lead2": "整理 ai/ay/a_e 的位置規律，回顧動物單字與 a/an，完成 3 題聽辨離堂卡。",
        "digital": "長母音自由拼音、動物翻卡、這那問答",
        "assessment": "形成性評量：A 長母音聽辨、動物詞彙辨識、Yes/No 完整回答。",
    },
    {
        "no": 7,
        "date": "115 年 6 月 2 日",
        "unit": "Action Studio 長母音 E/I 與動作表達",
        "source": "第三課",
        "language": "E 長母音 ee/ea/e_e、I 長母音 i_e/ie/igh、動作單字",
        "vocab": "sing, play, dance, jump, draw, run, sleep, write, read, cook, drink, eat",
        "patterns": "I can ___. / Can you ___? / It looks like ___.",
        "context": "情緒表達、身體動作與公共參與中的聆聽和輪流",
        "knowledge": "能辨識 E、I 長母音常見拼法，並理解 12 個動作動詞。",
        "practice": "能聽音選出 E/I 長母音，並用 I can... 表達一項能力。",
        "literacy": "能在動作任務中利用視覺、聲音與文字線索完成配對。",
        "lead1": "用 sea/bee/eve 與 bike/pie/light 建立 E、I 長母音字族；動作示範導入動詞。",
        "self": "在長母音拼音器自由替換字母、朗讀詞例；完成動作英文配對。",
        "group": "『無聲動作台』：學生抽卡演動作，隊友用 It looks like... 猜，再由表演者用 I can... 確認。",
        "share": "各組挑一個最有挑戰的長母音字與動作句，說明判斷線索。",
        "lead2": "以字族整理 E/I 拼法，對比易混淆音；用 5 題聽辨＋句型口說總結。",
        "digital": "長母音活動、動作圖卡、隨機挑人",
        "assessment": "形成性評量：E/I 長母音分類、動作單字辨識、I can... 句子產出。",
    },
    {
        "no": 8,
        "date": "115 年 6 月 9 日",
        "unit": "Can You Do It? O/OO 長短音與能力問答",
        "source": "第四課",
        "language": "O 長母音 oa/ow/o_e、OO 長短音、can/can’t/with",
        "vocab": "kick, talk, walk, close, clap, open; moon, food, book, cook",
        "patterns": "Can you ___? / Yes, I can. / No, I can’t. / I can ___ with ___.",
        "context": "均衡飲食、低碳生活與可實踐的健康行動",
        "knowledge": "能辨識 O 長母音與 OO 長短音，理解 can/can’t 的能力語意。",
        "practice": "能用 Can you...? 訪問同伴，並依真實情況開放式回答。",
        "literacy": "能操作任務卡記錄同伴回答，不把系統預設當成個人答案。",
        "lead1": "以 home/boat/snow 示範 O 長母音；以 food/book 比較 OO 長短音，再連結可做的健康行動。",
        "self": "完成 O/OO 聽辨與拼字分類；選一張行動卡，準備自己的 I can/I can’t 回答。",
        "group": "『能力任務所』：抽取人物、動作與夥伴卡，互問 Can you...?；回答者自行決定 Yes 或 No 並補充。",
        "share": "每組整理一項全組都會做與一項想學會的事，以 We can...／We want to... 分享。",
        "lead2": "整理 O/OO 音形關係與 can 問答語序，針對回答過短者示範補充資訊的方法。",
        "digital": "可不可以、長母音活動、圖卡配對",
        "assessment": "形成性評量：O/OO 聽辨、問句語序、依真實狀況回答與追問互動。",
    },
    {
        "no": 9,
        "date": "115 年 6 月 16 日",
        "unit": "Sports Challenge U 長母音與運動英語",
        "source": "第四課",
        "language": "U 長母音 u_e/ue/ui、運動單字、能力與喜好表達",
        "vocab": "soccer, basketball, baseball, dodgeball, volleyball, swim; cute, cube, blue, glue, fruit",
        "patterns": "Can you ___? / I can/can’t ___. / I like ___. / My favorite sport is ___.",
        "context": "運動後休息、補充水分與健康生活",
        "knowledge": "能辨識 U 長母音 /juː/、/uː/ 的常見拼法，並認讀 6 個運動單字。",
        "practice": "能詢問同伴會不會某項運動，並表達喜歡或最喜歡的運動。",
        "literacy": "能在洗牌翻卡活動中記憶位置、遵守回合並記錄配對結果。",
        "lead1": "比較 cute/blue/fruit 的 U 長母音，說明 /juː/ 與 /uː/；以動作線索導入 6 種運動。",
        "self": "完成 U 長母音分類與運動圖卡配對，練習 My favorite sport is...。",
        "group": "『運動配對聯賽』：卡片洗牌後輪流翻英文卡與中文卡，配對成功須說完整句才得分。",
        "share": "各組統計最常被選的運動，用簡單句分享結果與個人理由。",
        "lead2": "回顧 U 長母音、運動單字與 can/like 句型；進行聽音、配對、口說三關檢核。",
        "digital": "圖卡配對－運動組、長母音活動、隨機挑人",
        "assessment": "形成性評量：U 長母音分類、運動詞彙、能力問答與喜好句型完整度。",
    },
    {
        "no": 10,
        "date": "115 年 6 月 23 日",
        "unit": "English Mission Finale 長母音與生活溝通總任務",
        "source": "第四課／學期統整",
        "language": "A/E/I/O/U/OO 長母音總複習；this/that、Is it...?、can、主詞與 Be 動詞",
        "vocab": "本學期顏色、數字、文具、動物、動作與運動核心單字",
        "patterns": "What is this/that? / Is it a...? / Can you...? / I am... / It is...",
        "context": "溫馨時刻、數位安全、反詐與學期學習回顧",
        "knowledge": "能依拼字線索初步判斷長母音，並整合本學期核心詞彙與句型。",
        "practice": "能完成四站任務，從辨音、認字到問答與短句表達。",
        "literacy": "能辨識陌生連結與個資風險，知道遇到可疑訊息要停、想、問大人。",
        "lead1": "以聲音地圖統整 A/E/I/O/U/OO；示範四站任務規則與數位安全情境判斷。",
        "self": "完成個人長母音分類與核心單字自評，選出一項最有進步與一項仍想練習的能力。",
        "group": "輪流挑戰『辨音站、翻卡站、問答站、生活情境站』；每站由不同學生擔任記錄與確認角色。",
        "share": "各組展示一段包含至少 2 種句型的生活短對話，並分享一項合作策略。",
        "lead2": "完成後測與口語檢核，對照第一週起點資料，提供具體進步回饋並說明後續練習方向。",
        "digital": "長母音活動、圖卡配對、這那問答、可不可以、數位安全情境卡",
        "assessment": "總結性評量：辨音 30%、詞彙 25%、句型 25%、互動與任務參與 20%。",
    },
]


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = "w:" + edge
        node = tc_mar.find(qn(tag))
        if node is None:
            node = OxmlElement(tag)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="808080", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn("w:" + edge))
        if node is None:
            node = OxmlElement("w:" + edge)
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths, indent=TABLE_INDENT):
    assert sum(widths) == TABLE_WIDTH, (widths, sum(widths))
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def set_run_font(run, size=10.5, bold=False, color=INK, italic=False, latin=False):
    name = FONT_EN if latin else FONT
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=4, line=1.2, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def clear_cell(cell):
    cell.text = ""
    return cell.paragraphs[0]


def write_cell(cell, text, bold=False, size=9.6, color=INK, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = clear_cell(cell)
    style_paragraph(p, after=0, line=1.18, align=align)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold, color=color)
    return p


def add_label_value(cell, label, value, size=9.6):
    p = clear_cell(cell)
    style_paragraph(p, after=0, line=1.18)
    r = p.add_run(label)
    set_run_font(r, size=size, bold=True, color=TEAL_DARK)
    r = p.add_run(value)
    set_run_font(r, size=size, color=INK)


def add_heading(doc, text, level=1, kicker=None, page_break_before=False):
    if kicker:
        kp = doc.add_paragraph()
        style_paragraph(kp, before=3, after=2)
        kp.paragraph_format.page_break_before = page_break_before
        kr = kp.add_run(kicker.upper())
        set_run_font(kr, size=8.5, bold=True, color=TEAL, latin=True)
    p = doc.add_paragraph(style=f"Heading {level}")
    if page_break_before and not kicker:
        p.paragraph_format.page_break_before = True
    r = p.add_run(text)
    return p


def add_body(doc, text, bold_label=None, after=5):
    p = doc.add_paragraph()
    style_paragraph(p, after=after, line=1.25)
    if bold_label:
        r = p.add_run(bold_label)
        set_run_font(r, size=10.5, bold=True, color=TEAL_DARK)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=INK)
    return p


def add_section_band(doc, text):
    t = doc.add_table(rows=1, cols=1)
    set_table_geometry(t, [TABLE_WIDTH])
    set_table_borders(t, color=TEAL, size="4")
    set_cell_shading(t.cell(0, 0), PALE_TEAL)
    write_cell(t.cell(0, 0), text, bold=True, size=11.5, color=TEAL_DARK, align=WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    set_run_font(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)
    run2 = paragraph.add_run(" 頁")
    set_run_font(run2, size=8.5, color=MID_GRAY)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.78)
section.bottom_margin = Inches(0.72)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.header_distance = Inches(0.35)
section.footer_distance = Inches(0.35)
section.different_first_page_header_footer = True

styles = doc.styles
normal = styles["Normal"]
normal.font.name = FONT
normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(INK)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in (
    ("Title", 26, INK, 0, 8),
    ("Subtitle", 13, MID_GRAY, 0, 8),
    ("Heading 1", 17, TEAL_DARK, 10, 6),
    ("Heading 2", 13, TEAL_DARK, 8, 5),
    ("Heading 3", 11.5, INK, 6, 3),
):
    st = styles[name]
    st.font.name = FONT
    st._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    st._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    st._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = name != "Subtitle"
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("臺東大學 114 學年第 2 學期 AiDi+ 大學伴教學單元活動設計")
set_run_font(hr, size=8.5, bold=True, color=MID_GRAY)

footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

p = doc.add_paragraph()
style_paragraph(p, after=12, align=WD_ALIGN_PARAGRAPH.CENTER)
r = p.add_run("臺東大學 114 學年第 2 學期\nAiDi+ 大學伴教學單元活動設計")
set_run_font(r, size=18, bold=True, color=INK)

cover = doc.add_table(rows=3, cols=2)
set_table_geometry(cover, [4680, 4680])
set_table_borders(cover, color="808080", size="6")
cover_data = [
    (("領域／科目　", "英語文"), ("設計者　", "顏子薰")),
    (("實施年級　", "國小三年級"), ("教學節次　", "10 次，共 20 節")),
    (("教材依據　", "何嘉仁第二冊、因材網"), ("課程期間　", "115 年 4 月至 6 月")),
]
for i, row_data in enumerate(cover_data):
    for j, (label, value) in enumerate(row_data):
        set_cell_shading(cover.cell(i, j), PALE_GRAY if i % 2 else PALE_GOLD)
        add_label_value(cover.cell(i, j), label, value)

add_heading(doc, "一、課程整體設計", 1)
add_body(
    doc,
    "本課程依據 10 份實際授課教材，將本學期內容重整為『語音覺識－核心詞彙－句型溝通－生活應用』的漸進路徑。每次課程以學生熟悉的生活經驗、屏東與臺東在地情境作為語意支架，再透過自編數位活動進行可視化操作、即時回饋與合作任務。課程不以大量背誦為主，而是讓學生在聽、看、選、排、說的循環中，逐步建立音形連結與開口信心。",
)

add_section_band(doc, "設計依據與核心素養")
foundation = doc.add_table(rows=4, cols=2)
set_table_geometry(foundation, [2050, 7310])
set_table_borders(foundation)
foundation_data = [
    ("學習表現", "1-II-7 能聽懂課堂中所學的字詞。\n2-II-3 能說出課堂中所學的字詞。\n6-II-2 積極參與各種課堂練習活動。"),
    ("學習內容", "Aa-II-2 印刷體大小寫字母的辨識及書寫。\nAb-II-1 子音、母音及其組合。\nAc-II-2 簡易的生活用語。"),
    ("核心素養", "英-E-A1 建立專注與良好學習習慣，運用基本策略強化英語能力。\n英-E-A2 理解簡易英語訊息，運用基本邏輯思考提升學習效能。\n英-E-B1 在引導下運用字詞與句型進行簡易日常溝通。"),
    ("議題融入", "戶外教育、生命教育、多元文化、海洋教育、資訊與數位安全。議題作為英語溝通情境，不取代語言學習主軸。"),
]
for i, (label, value) in enumerate(foundation_data):
    set_cell_shading(foundation.cell(i, 0), PALE_GOLD)
    write_cell(foundation.cell(i, 0), label, bold=True, size=9.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(foundation.cell(i, 1), value, size=9.4)

add_heading(doc, "學期總目標", 2)
goals = doc.add_table(rows=4, cols=2)
set_table_geometry(goals, [2050, 7310])
set_table_borders(goals)
goal_data = [
    ("語音覺識", "辨識短母音 a/e/i/o/u，以及 A/E/I/O/U/OO 常見長短音拼法，能在提示下進行拆音與合音。"),
    ("詞彙理解", "認讀並運用顏色、數字、文具、動物、動作與運動等生活核心詞彙。"),
    ("句型溝通", "使用自我介紹、What is this/that?、Is it...?、Can you...? 等句型完成簡短互動。"),
    ("學習素養", "能輪流參與、接受多元作答方式、閱讀數位回饋並依提示修正答案。"),
]
for i, (label, value) in enumerate(goal_data):
    set_cell_shading(goals.cell(i, 0), PALE_TEAL if i % 2 == 0 else PALE_GOLD)
    write_cell(goals.cell(i, 0), label, bold=True, size=9.8, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(goals.cell(i, 1), value, size=9.4)

add_heading(doc, "適性參與原則", 2)
add_body(
    doc,
    "四位學生皆可依當下狀態選擇口說、指認、拖曳、字卡排序或點頭／搖頭等回應方式；教師以清楚單一任務、充足等待時間、可重複操作與具體正向回饋支持參與。評量重點放在相較自身起點的進步，不以音量、反應速度或一次作答結果作為唯一判準。",
)

add_heading(doc, "二、學期課程地圖", 1, page_break_before=True)
add_body(doc, "每次教學為 2 節、約 80 分鐘；課間休息可依學校作息彈性安排。")

overview = doc.add_table(rows=1, cols=5)
set_table_geometry(overview, [760, 1500, 2400, 2300, 2400])
set_table_borders(overview, color="808080", size="6")
headers = ["次數", "日期", "單元", "語言焦點", "主要數位活動"]
for i, h in enumerate(headers):
    set_cell_shading(overview.cell(0, i), TEAL)
    write_cell(overview.cell(0, i), h, bold=True, size=9.0, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(overview.rows[0])
for s in sessions:
    row = overview.add_row()
    vals = [str(s["no"]), s["date"].replace("115 年 ", "").replace(" 月 ", "/").replace(" 日", ""), s["unit"], s["language"], s["digital"]]
    for i, val in enumerate(vals):
        if s["no"] % 2 == 0:
            set_cell_shading(row.cells[i], PALE_GRAY)
        write_cell(row.cells[i], val, size=8.2, align=WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(overview, [760, 1500, 2400, 2300, 2400])

add_heading(doc, "共同教學節奏", 2, page_break_before=True)
rhythm = doc.add_table(rows=5, cols=3)
set_table_geometry(rhythm, [2050, 1300, 6010])
set_table_borders(rhythm)
rhythm_data = [
    ("教師導學", "20 分", "建立情境、示範語音或句型、說明任務成功條件"),
    ("學生自學", "15 分", "個別聽辨、拖曳、配對、跟讀或自我修正"),
    ("組內共學", "20 分", "四人輪流角色、合作解題與完整句練習"),
    ("組間互學", "10 分", "展示策略、同儕判斷、說明理由或補充答案"),
    ("教師導學", "15 分", "釐清迷思、統整語言規則、離堂檢核與回饋"),
]
for i, vals in enumerate(rhythm_data):
    for j, val in enumerate(vals):
        if j == 0:
            set_cell_shading(rhythm.cell(i, j), PALE_GOLD)
        elif j == 1:
            set_cell_shading(rhythm.cell(i, j), PALE_TEAL)
        write_cell(rhythm.cell(i, j), val, bold=j < 2, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT)

add_heading(doc, "共通數位活動流程", 2)
common_digital = doc.add_table(rows=4, cols=3)
set_table_geometry(common_digital, [1550, 2350, 5460])
set_table_borders(common_digital)
common_digital_data = [
    ("1 準備", "確認介面與任務", "教師先完成一次操作，確認拖曳、觸控、聲音、名字與學生回合設定。"),
    ("2 示範", "說出思考步驟", "先示範『看線索－選答案－讀回饋－需要時修正』，再讓學生操作。"),
    ("3 操作", "維持流暢回合", "避免每次點選都重整頁面；同一任務中保留學生已完成的狀態與分數。"),
    ("4 回饋", "留下學習證據", "記錄學生獨立完成、提示後完成與偏好的回應方式，不只記錄對錯。"),
]
for i, vals in enumerate(common_digital_data):
    for j, val in enumerate(vals):
        if j == 0:
            set_cell_shading(common_digital.cell(i, j), PALE_GOLD)
        elif j == 1:
            set_cell_shading(common_digital.cell(i, j), PALE_TEAL)
        write_cell(common_digital.cell(i, j), val, bold=j < 2, size=9.0, align=WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT)

add_heading(doc, "每堂課通用檢核", 2)
add_body(doc, "課前確認網站與備援素材；課中確認每位學生至少有一次主動操作及一次語言回應；課後記錄一項亮點、一項待練習概念與下一次調整。")

# Session detail pages.
for s in sessions:
    add_heading(doc, f"第 {s['no']} 次教學活動設計：{s['unit']}", 1, page_break_before=True)

    meta = doc.add_table(rows=3, cols=2)
    set_table_geometry(meta, [4680, 4680])
    set_table_borders(meta)
    meta_rows = [
        (("日期　", s["date"]), ("教材進度　", s["source"])),
        (("語言焦點　", s["language"]), ("情境脈絡　", s["context"])),
        (("核心字詞　", s["vocab"]), ("核心句型　", s["patterns"])),
    ]
    for i, pair in enumerate(meta_rows):
        for j, (label, value) in enumerate(pair):
            set_cell_shading(meta.cell(i, j), PALE_TEAL if (i + j) % 2 == 0 else PALE_GRAY)
            add_label_value(meta.cell(i, j), label, value, size=8.9)

    add_heading(doc, "學習目標", 2)
    objective = doc.add_table(rows=3, cols=2)
    set_table_geometry(objective, [2050, 7310])
    set_table_borders(objective)
    for i, (label, value) in enumerate(
        [
            ("學科／知識能力", s["knowledge"]),
            ("溝通／任務表現", s["practice"]),
            ("自主學習與數位素養", s["literacy"]),
        ]
    ):
        set_cell_shading(objective.cell(i, 0), PALE_GOLD)
        write_cell(objective.cell(i, 0), label, bold=True, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(objective.cell(i, 1), value, size=9.2)

    add_heading(doc, "課堂四學流程（80 分鐘）", 2)
    flow = doc.add_table(rows=1, cols=4)
    set_table_geometry(flow, [1450, 850, 5160, 1900])
    set_table_borders(flow, color="808080", size="6")
    for i, h in enumerate(["階段", "時間", "學習活動／任務", "資源／評量"]):
        set_cell_shading(flow.cell(0, i), TEAL)
        write_cell(flow.cell(0, i), h, bold=True, size=8.9, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
    set_repeat_table_header(flow.rows[0])
    flow_data = [
        ("教師導學", "20", s["lead1"], "簡報、實物／口頭提問"),
        ("學生自學", "15", s["self"], "個別操作／任務完成度"),
        ("組內共學", "20", s["group"], "角色卡／互動觀察"),
        ("組間互學", "10", s["share"], "展示／同儕回饋"),
        ("教師導學", "15", s["lead2"], "離堂卡／即時回饋"),
    ]
    for i, vals in enumerate(flow_data):
        row = flow.add_row()
        for j, value in enumerate(vals):
            if i % 2 == 1:
                set_cell_shading(row.cells[j], PALE_GRAY)
            if j == 0:
                set_cell_shading(row.cells[j], PALE_GOLD)
            write_cell(row.cells[j], value, bold=j < 2, size=8.6, align=WD_ALIGN_PARAGRAPH.CENTER if j < 2 else WD_ALIGN_PARAGRAPH.LEFT)
        set_table_geometry(flow, [1450, 850, 5160, 1900])

    add_heading(doc, "數位活動與評量證據", 2)
    evidence = doc.add_table(rows=2, cols=2)
    set_table_geometry(evidence, [2050, 7310])
    set_table_borders(evidence)
    for i, (label, value) in enumerate([("AI／數位活動", s["digital"]), ("評量方式", s["assessment"])]):
        set_cell_shading(evidence.cell(i, 0), PALE_TEAL if i == 0 else PALE_GOLD)
        write_cell(evidence.cell(i, 0), label, bold=True, size=9.2, align=WD_ALIGN_PARAGRAPH.CENTER)
        write_cell(evidence.cell(i, 1), value, size=9.2)

# Semester assessment appendix.
add_heading(doc, "三、學期評量設計", 1, page_break_before=True)
add_body(
    doc,
    "採診斷性、形成性與總結性評量並行。每次活動記錄學生的『獨立完成、提示後完成、同伴協助完成、尚待練習』狀態；學期末以相同或相近任務對照起點，不以單次分數取代學習歷程。",
)

rubric = doc.add_table(rows=1, cols=5)
set_table_geometry(rubric, [1700, 1915, 1915, 1915, 1915])
set_table_borders(rubric, color="808080", size="6")
for i, h in enumerate(["向度", "4 穩定運用", "3 大致達成", "2 提示完成", "1 持續建立"]):
    set_cell_shading(rubric.cell(0, i), TEAL)
    write_cell(rubric.cell(0, i), h, bold=True, size=8.8, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(rubric.rows[0])
rubric_data = [
    ("語音覺識", "能依音形線索辨音並嘗試拼讀新詞", "常見詞辨音正確，偶有混淆", "在選項或口型提示下完成", "仍需逐音示範與反覆比對"),
    ("核心詞彙", "能快速辨認並在句中使用", "多數詞彙可正確辨認", "看圖或首音提示後作答", "需以配對與重複接觸建立"),
    ("句型表達", "能完整、適切且自主回應", "能用目標句型完成回答", "依句框或字卡完成", "以指認、單字或模仿參與"),
    ("互動參與", "能輪流、傾聽、回應並協助同伴", "能遵守回合並完成任務", "提醒後能回到任務", "需要較多等待與明確單一步驟"),
    ("數位素養", "能自主操作、讀回饋並修正", "能完成操作與基本修正", "在示範或同伴協助下完成", "仍需逐步引導與介面熟悉"),
]
for r_idx, vals in enumerate(rubric_data):
    row = rubric.add_row()
    for c_idx, value in enumerate(vals):
        if r_idx % 2 == 1:
            set_cell_shading(row.cells[c_idx], PALE_GRAY)
        if c_idx == 0:
            set_cell_shading(row.cells[c_idx], PALE_GOLD)
        write_cell(row.cells[c_idx], value, bold=c_idx == 0, size=8.25, align=WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(rubric, [1700, 1915, 1915, 1915, 1915])

add_heading(doc, "評量證據配置", 2)
evidence_grid = doc.add_table(rows=4, cols=3)
set_table_geometry(evidence_grid, [1900, 3730, 3730])
set_table_borders(evidence_grid)
evidence_rows = [
    ("時間", "評量任務", "紀錄重點"),
    ("學期初", "字母、數字、生活用語、常見單字與短句診斷", "正確度、回應方式、等待時間、可接受提示"),
    ("每次課程", "離堂卡、口說任務、數位操作、合作活動", "音形、詞彙、句型、互動與修正歷程"),
    ("學期末", "長母音分類、核心詞彙、情境問答與短對話", "與起點對照的進步、可遷移能力與後續需求"),
]
for i, vals in enumerate(evidence_rows):
    for j, val in enumerate(vals):
        if i == 0:
            set_cell_shading(evidence_grid.cell(i, j), TEAL)
            write_cell(evidence_grid.cell(i, j), val, bold=True, size=9, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
        else:
            if j == 0:
                set_cell_shading(evidence_grid.cell(i, j), PALE_GOLD)
            write_cell(evidence_grid.cell(i, j), val, bold=j == 0, size=9)

add_heading(doc, "四、教學資源與課前檢核", 1, page_break_before=True)
resource = doc.add_table(rows=5, cols=2)
set_table_geometry(resource, [2050, 7310])
set_table_borders(resource)
resource_data = [
    ("教材來源", "何嘉仁第二冊、因材網、textbook/English 之 10 份授課簡報、自編 AiDi+ 互動活動。"),
    ("教學設備", "桌面式電腦、投影或共享畫面、滑鼠／觸控裝置、喇叭、實物與圖卡。"),
    ("課前確認", "網頁可開啟、拖曳與觸控正常、聲音可播放、學生名字設定正確、活動不會因單次操作整頁重整。"),
    ("備援方案", "若網路不穩，使用紙本字卡、顏色卡、動作卡與白板句框維持相同任務目標。"),
    ("資料與安全", "不公開學生個人紀錄；使用作品或截圖前先確認用途；數位安全情境只蒐集必要資訊。"),
]
for i, (label, value) in enumerate(resource_data):
    set_cell_shading(resource.cell(i, 0), PALE_GOLD if i % 2 == 0 else PALE_TEAL)
    write_cell(resource.cell(i, 0), label, bold=True, size=9.4, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(resource.cell(i, 1), value, size=9.4)

add_heading(doc, "教師課後紀錄欄", 2)
reflection = doc.add_table(rows=5, cols=2)
set_table_geometry(reflection, [2050, 7310])
set_table_borders(reflection)
for i, label in enumerate(["本次完成", "學生亮點", "待釐清概念", "下次調整", "個別狀態"]):
    set_cell_shading(reflection.cell(i, 0), PALE_TEAL)
    write_cell(reflection.cell(i, 0), label, bold=True, size=9.4, align=WD_ALIGN_PARAGRAPH.CENTER)
    write_cell(reflection.cell(i, 1), "\n", size=9.4)

add_heading(doc, "教材對照", 2, page_break_before=True)
materials = doc.add_table(rows=1, cols=3)
set_table_geometry(materials, [900, 2200, 6260])
set_table_borders(materials, color="808080")
for i, h in enumerate(["次數", "日期", "參考教材檔案"]):
    set_cell_shading(materials.cell(0, i), TEAL)
    write_cell(materials.cell(0, i), h, bold=True, size=8.8, color=WHITE, align=WD_ALIGN_PARAGRAPH.CENTER)
set_repeat_table_header(materials.rows[0])
material_files = [
    "仁和國小英文Res3_20260421.pdf",
    "仁和國小英文Res3_20260428.pdf",
    "仁和國小英文Res3_20260505.pdf",
    "仁和國小英文Res3_20260512.pdf",
    "仁和國小英文Res3_20260519.pdf",
    "仁和國小英文Res3_20260526.pdf",
    "仁和國小英文Res3_20260602.pdf",
    "仁和國小英文Res3_20260609.pdf",
    "仁和國小英文Res3_20260616.pdf",
    "仁和國小英文Res3_20260623.pdf",
]
for s, filename in zip(sessions, material_files):
    row = materials.add_row()
    for i, val in enumerate([str(s["no"]), s["date"], filename]):
        if s["no"] % 2 == 0:
            set_cell_shading(row.cells[i], PALE_GRAY)
        write_cell(row.cells[i], val, size=8.6, align=WD_ALIGN_PARAGRAPH.CENTER if i < 2 else WD_ALIGN_PARAGRAPH.LEFT)
    set_table_geometry(materials, [900, 2200, 6260])
add_body(doc, "以上教材依實際授課日期排序；各次教案已擷取語音、詞彙、句型、生活情境與數位活動重點，重新組織為可連續實施的一學期課程。", after=2)

# Representative screenshots from the teaching website.
screenshots = [
    ("tmp/web_activity_crops/phonics.png", "母音練習：長短音比較、長母音與自由拼音"),
    ("tmp/web_activity_crops/color-mix.png", "顏色小遊戲：聽顏色、辨識色塊與顏色混合"),
    ("tmp/web_activity_crops/this-that.png", "這那問答：以翻卡搜查任務練習 this / that 與 Yes / No"),
    ("tmp/web_activity_crops/can-do.png", "可不可以：抽取動作與夥伴卡，進行開放式 Can you...? 問答"),
    ("tmp/web_activity_crops/card-match.png", "圖卡配對：依教材主題切換卡組，進行英中配對與記憶挑戰"),
    ("tmp/web_activity_crops/random-picker.png", "隨機挑人：用多種動畫公平安排回合與角色"),
]
for page_index, (path, caption) in enumerate(screenshots):
    add_heading(doc, "五、數位教學活動畫面", 1, page_break_before=True)
    add_body(doc, f"圖 {page_index + 1}　{caption}", after=7)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run()
    inline_shape = run.add_picture(path, width=Inches(6.15))
    inline_shape._inline.docPr.set("descr", "數位教學網站截圖：" + caption)
    inline_shape._inline.docPr.set("title", caption)

# Document metadata.
props = doc.core_properties
props.title = "AiDi+ 英文教學單元活動設計（一學期完整版）"
props.subject = "國小三年級英語，一學期 10 次、20 節教學活動設計"
props.author = "顏子薰"
props.keywords = "AiDi+, 數位學伴, 國小三年級, 英文教案, 長短母音, 數位活動"

doc.save(OUTPUT)
print(OUTPUT.resolve())
